# -*- coding: utf-8 -*-
"""사내 공유용 소개 PPT 생성 (2026-09-13).

내용의 원본은 docs/사내공유_메일_20260913.md 이고, 이 스크립트는 **그 메일과 같은 순서·같은 말**로
슬라이드를 만든다(메일 1~5항 → 본문 13장 + 부록 2장). 메일과 자료가 다른 이야기를 하면 안 된다.
시스템이 바뀌면 메일 원본 → 이 스크립트 순으로 고친다. AI 호출 없음.

실행: C:/Users/SKTelecom/AppData/Local/Programs/Python/Python312/python.exe make_intro_deck.py
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

W, H   = Inches(13.333), Inches(7.5)
INK    = RGBColor(0x1A, 0x1A, 0x1A)
MUTED  = RGBColor(0x5F, 0x66, 0x73)
ACCENT = RGBColor(0xEA, 0x00, 0x2C)      # SKT 레드
NAVY   = RGBColor(0x0F, 0x1B, 0x2E)
PANEL  = RGBColor(0xF4, 0xF5, 0xF7)
LINE   = RGBColor(0xDF, 0xE3, 0xE9)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DIM    = RGBColor(0xB0, 0xB8, 0xC4)
FONT   = '맑은 고딕'

MARGIN = Inches(0.75)
BODY_W = W - 2 * MARGIN


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, round_=True):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    if round_:
        sh.adjustments[0] = 0.035
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def bar(slide, x, y, w, h, fill):
    return rect(slide, x, y, w, h, fill, None, round_=False)


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=Pt(4), line_spacing=1.2):
    """runs: [(문자열, pt, 굵게, 색)] — 항목 하나가 문단 하나."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (s, size, bold, color) in enumerate(runs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = space_after
        para.line_spacing = line_spacing
        r = para.add_run(); r.text = s
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = FONT
    return tb


def head(slide, title, sub=None, idx=None, tag=None):
    bar(slide, MARGIN, Inches(0.6), Inches(0.085), Inches(0.44), ACCENT)
    text(slide, MARGIN + Inches(0.22), Inches(0.56), Inches(9.5), Inches(0.5),
         [(title, 26, True, INK)])
    if tag:
        text(slide, W - MARGIN - Inches(3.0), Inches(0.66), Inches(3.0), Inches(0.3),
             [(tag, 11, False, DIM)], align=PP_ALIGN.RIGHT)
    if sub:
        text(slide, MARGIN + Inches(0.22), Inches(1.11), BODY_W - Inches(0.5), Inches(0.4),
             [(sub, 12.5, False, MUTED)])
    if idx is not None:
        text(slide, W - MARGIN - Inches(1.0), H - Inches(0.6), Inches(1.0), Inches(0.3),
             [(str(idx), 10, False, DIM)], align=PP_ALIGN.RIGHT)


def cards(slide, items, top, cols=3, card_h=Inches(2.0), gap=Inches(0.22),
          head_size=14.5, body_size=11.5):
    """items: [(제목, 본문)] 또는 [(번호, 제목, 본문)]"""
    cw = (BODY_W - gap * (cols - 1)) / cols
    for i, it in enumerate(items):
        row, col = divmod(i, cols)
        x = MARGIN + col * (cw + gap)
        y = top + row * (card_h + gap)
        rect(slide, x, y, cw, card_h, PANEL, LINE)
        pad = Inches(0.26)
        if len(it) == 3:
            num, title, body = it
            text(slide, x + pad, y + Inches(0.2), cw - 2 * pad, Inches(0.3),
                 [(num, 13, True, ACCENT)])
            hy = y + Inches(0.55)
        else:
            title, body = it
            hy = y + Inches(0.24)
        text(slide, x + pad, hy, cw - 2 * pad, Inches(0.4), [(title, head_size, True, INK)])
        text(slide, x + pad, hy + Inches(0.4), cw - 2 * pad, card_h - Inches(1.0),
             [(body, body_size, False, MUTED)], line_spacing=1.32)


def rows(slide, data, top, left_w=Inches(2.6), row_h=Inches(0.62),
         left_size=12, right_size=11.5):
    y = top
    for i, (a, b) in enumerate(data):
        if i % 2 == 0:
            bar(slide, MARGIN, y, BODY_W, row_h, PANEL)
        text(slide, MARGIN + Inches(0.24), y + Inches(0.13), left_w - Inches(0.32), row_h,
             [(a, left_size, True, INK)], line_spacing=1.2)
        text(slide, MARGIN + left_w, y + Inches(0.13), BODY_W - left_w - Inches(0.3), row_h,
             [(b, right_size, False, MUTED)], line_spacing=1.28)
        y += row_h
    return y


def note(slide, s, y=None):
    y = y or (H - Inches(1.02))
    bar(slide, MARGIN, y, Inches(0.05), Inches(0.4), ACCENT)
    text(slide, MARGIN + Inches(0.18), y + Inches(0.04), BODY_W - Inches(0.4), Inches(0.4),
         [(s, 11.5, False, INK)], line_spacing=1.25)


# ══════════════════════════════════════════════════════════════
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # ── 01 표지 ───────────────────────────────────────────────
    s = blank(prs)
    bar(s, 0, 0, W, H, NAVY)
    bar(s, MARGIN, Inches(2.5), Inches(0.5), Inches(0.07), ACCENT)
    text(s, MARGIN, Inches(2.8), Inches(11.0), Inches(1.2),
         [('전파·통신 정책 모니터링 · AI 자문 시스템', 38, True, WHITE)])
    text(s, MARGIN, Inches(3.9), Inches(11.0), Inches(0.9),
         [('법령 관계도 · 자동 수집 · 근거를 대조하는 AI 자문', 15, False, DIM)])
    text(s, MARGIN, Inches(5.85), Inches(11.0), Inches(1.0),
         [('SKT Comm.센터 기술정책팀   ·   2026. 9.', 12.5, False, RGBColor(0x7E, 0x88, 0x99)),
          ('radio-policy.gitlab.io      ·      t.me/radio_policy_law_ai_bot', 12.5, True, WHITE)],
         space_after=Pt(7))

    # ── 02 Agent 개요 (메일 1항) ──────────────────────────────
    s = blank(prs); head(s, 'Agent 개요', None, 2, '메일 1항')
    rect(s, MARGIN, Inches(1.35), BODY_W, Inches(1.25), PANEL, LINE)
    text(s, MARGIN + Inches(0.32), Inches(1.58), BODY_W - Inches(0.64), Inches(0.9),
         [('대외 업무를 하다 보면 법률과 관련해 검토할 일이 생깁니다.', 14, True, INK),
          ('AI에 맡기면 없는 조문을 지어내고, 직접 찾으면 시행령·시행규칙·고시·별표까지 타고 '
           '내려가야 합니다. 이 두 가지를 풀려고 만들었습니다.', 12.5, False, MUTED)],
         space_after=Pt(6), line_spacing=1.3)
    cards(s, [
        ('01', '법령 관계도',
         '법률에서 시행령·시행규칙·고시·별표로 이어지는 위임 관계와 조문 사이 인용 관계를 '
         '자동으로 추출해 지도로 보여 줍니다.'),
        ('02', '자동 수집',
         '정부 보도자료, 국내외 규제기관 동향, 국회 회의록과 법안, 정책 뉴스를 매일 모으고, '
         '법령이 개정되면 바뀐 조문을 비교해 알려 줍니다.'),
        ('03', 'AI 자문',
         '관련 법령·고시 원문에서 근거를 찾아 답하고, 인용한 조문이 원문과 맞는지 기계가 대조해 '
         '[원문 확인됨]을 붙입니다.'),
    ], Inches(2.95), cols=3, card_h=Inches(2.15))
    note(s, '설치할 것도, 따로 드는 비용도 없습니다 — 웹 주소와 텔레그램 링크가 전부입니다.')

    # ── 03 법령 관계도 ────────────────────────────────────────
    s = blank(prs); head(s, '① 법령 관계도', '조문 하나가 어디에 물려 있는지 지도로 봅니다', 3)
    rows(s, [
        ('무엇이 문제였나',
         '전파법 조문 하나를 확인하려면 시행령 → 시행규칙 → 고시 → 별표까지 타고 내려가야 하고,\n'
         '어디까지 봐야 다 본 것인지 알 수 없습니다'),
        ('위임 관계',
         '법제처 3단비교 정본에서 법률↔시행령·시행규칙 조문 대응을 가져오고,\n'
         '고시는 제1조(목적)의 근거 문구를 역추출해 상위 법령에 잇습니다'),
        ('인용 관계',
         '조문이 다른 조문을 인용한 선을 그려, 벌칙·검사 조항처럼 검색어로는 닿지 않는 곳까지 보여 줍니다'),
        ('주제별 보기',
         '"무선통신시설 공동이용" 같은 주제를 고르면 그 주제에 걸린 법령·고시만 모아 봅니다.\n'
         '조문 단위로도 펼칠 수 있고, 노드를 누르면 원문으로 바로 갑니다'),
    ], Inches(1.9), left_w=Inches(2.5), row_h=Inches(1.02))
    note(s, 'AI가 새로 제안한 연결은 바로 반영하지 않고, 운영자가 원문을 확인해 승인한 것만 지도에 올립니다.')

    # ── 04 자동 수집 ──────────────────────────────────────────
    s = blank(prs); head(s, '② 자동 수집', '정해진 시각에 정해진 일이 돕니다 — 사람이 챙길 일이 없습니다', 4)
    cards(s, [
        ('매시간', '정책 뉴스 수집 → 관련성 판정 → 중요도 분류.\n중요 건은 시간을 가리지 않고 알립니다'),
        ('06:00', '모닝 브리핑 발송.\n어제~오늘 핵심 + 업무 영향 분석'),
        ('11:00', '법령·고시 개정 감지.\n바뀐 조문을 신·구 대조해 보여 줍니다'),
        ('17:00', '국회·부처 입법예고 수집,\n과방위 회의록 등재·요약'),
    ], Inches(1.9), cols=4, card_h=Inches(1.55), head_size=15, body_size=11)
    text(s, MARGIN, Inches(3.62), BODY_W, Inches(0.35), [('어디에서 모으나', 13.5, True, INK)])
    rows(s, [
        ('정부', '과학기술정보통신부 · 국립전파연구원 · 중앙전파관리소 · 방송미디어통신위원회 · ETRI · KISDI'),
        ('법령 · 국회', '법제처 국가법령정보 · 열린국회정보(법안·입법예고) · 국회 회의록시스템'),
        ('해외 규제기관', 'FCC(미국) · Ofcom(영국) · BEREC(EU) · 총무성(일본) · ITU'),
        ('뉴스', '네이버 검색 API + 언론사 직접 수집 (수집 키워드는 부록 A)'),
    ], Inches(4.05), left_w=Inches(2.3), row_h=Inches(0.58), left_size=11.5, right_size=11)
    note(s, '자동 실행은 두 벌로 돌립니다 — 한쪽이 멈춰도 다른 쪽이 대신하고, 멈추면 감시 장치가 알립니다.')

    # ── 05 AI 자문 ────────────────────────────────────────────
    s = blank(prs); head(s, '③ AI 자문', '답보다 근거가 먼저입니다', 5)
    cards(s, [
        ('근거를 찾아 답합니다',
         '질문에 맞는 법령·고시 조문, 정부 보도자료, 수집한 뉴스를 함께 찾아 근거로 삼습니다.\n\n'
         '조문이 값을 별표로 넘기면 그 별표까지 함께 읽어 옵니다.'),
        ('근거를 기계가 대조합니다',
         '인용한 조문이 실제로 있는 내용인지 대조해 [원문 확인됨]을 붙입니다.\n\n'
         '못 찾으면 「원문 미확인」, 원문과 다르게 설명했으면 「확인 필요」로 바꿔 돌려줍니다 — '
         'AI가 스스로 붙이는 표시가 아닙니다.'),
    ], Inches(1.95), cols=2, card_h=Inches(2.35), head_size=16)
    text(s, MARGIN, Inches(4.35), BODY_W, Inches(0.3),
         [('답변에 붙는 표시는 세 가지입니다', 13, True, INK)])
    rows(s, [
        ('[원문 확인됨]', '인용한 조문이 실제로 있고, 설명도 원문과 어긋나지 않습니다'),
        ('⚠ 원문 미확인', '검색 결과에 해당 조문이 없습니다 — 그대로 믿지 마십시오'),
        ('⚠ 확인 필요', '조문은 있으나 원문과 다르게 설명했습니다'),
    ], Inches(4.72), left_w=Inches(2.6), row_h=Inches(0.5), left_size=11.5, right_size=11)
    note(s, '법안은 아직 법이 아니므로 자문의 근거로 쓰지 않습니다 — 「국회 동향」으로만 덧붙입니다.')

    # ── 06 대시보드 지도 (메일 2항) ───────────────────────────
    s = blank(prs); head(s, '대시보드 — 열람은 로그인 없이',
                         'radio-policy.gitlab.io  ·  PC와 휴대폰 모두 지원', 6, '메일 2항')
    rows(s, [
        ('법령 · 모니터링', '법령·고시 조문 검색 · 관계도 · 개정 조문 비교 · 용어  |  정책 뉴스(자동 분류, 60일) · 아침 브리핑 · 정부·해외 보도자료'),
        ('이슈맵', 'AI가 중요한 흐름을 이슈로 제안하고, 운영자가 승인하면 관련 뉴스·보도자료·법령·이해관계자를 연대기로 정리'),
        ('국회', '법안 발의·처리 경과와 입법예고 · 과방위 회의록(2016년 20대 국회부터) · 인물 243명의 발언 이력과 쟁점별 입장 요약'),
        ('AI 자문', '질문을 입력하면 근거 조문과 함께 답변이 나옵니다 — 로그인·승인이 필요한 유일한 기능입니다'),
    ], Inches(2.15), left_w=Inches(2.2), row_h=Inches(0.95), left_size=13.5)
    note(s, '링크를 눌러 바로 확인해 보셔도 됩니다 — 로그인이 필요한 것은 AI 자문뿐입니다.')

    # ── 07 법령 탭 ────────────────────────────────────────────
    s = blank(prs); head(s, '대시보드 ① 법령', '원문이 정답인 자리에는 AI 해설을 붙이지 않습니다', 7)
    rows(s, [
        ('조문 검색', '법령·고시 원문을 조문 단위로 검색. 「제N조」 원문과 시행일을 함께 표시합니다'),
        ('개정 비교', '법이 바뀌면 어느 조문이 어떻게 바뀌었는지 신·구 대조로 보여 줍니다.\n'
                      '시행예정 법령도 미리 비교해 둡니다'),
        ('법적 용어', '조문에 정의된 용어를 원문 그대로. 출처(법령명·조·호·시행일)까지 붙습니다.\n'
                      'AI 해설을 붙이지 않습니다 — 법률 용어는 조문이 정답입니다'),
        ('기술 용어', '뉴스에서 새 기술 용어를 자동으로 뽑아 풀이합니다'),
    ], Inches(1.9), left_w=Inches(2.2), row_h=Inches(1.0))

    # ── 08 모니터링 탭 ────────────────────────────────────────
    s = blank(prs); head(s, '대시보드 ② 모니터링', '매시간 모으고, 아침에 한 통으로 정리해 보냅니다', 8)
    rows(s, [
        ('정책 뉴스', '중요도 3단계 자동 분류 — 즉시대응 / 금주검토 / 동향파악.\n'
                      '60일 보관하며, 남겨 둘 기사는 잠가서 계속 보관합니다'),
        ('모닝 브리핑', '매일 06:00 발송. 어제~오늘 핵심 + 업무 영향 분석 + 새로 나온 용어 풀이.\n'
                        '지난 브리핑이 날짜별로 전부 남아 있어 되짚어 볼 수 있습니다'),
        ('정부 공지', '과기정통부·전파연구원·전파관리소·방미통위의 공고와 보도자료 전문'),
        ('해외 동향', 'FCC·Ofcom·BEREC·일본 총무성·ITU의 규제 동향'),
    ], Inches(1.9), left_w=Inches(2.2), row_h=Inches(1.0))
    note(s, '중요도 기준은 팀마다 달라야 합니다 — 13장을 봐 주십시오.')

    # ── 09 이슈맵 탭 ──────────────────────────────────────────
    s = blank(prs); head(s, '대시보드 ③ 이슈맵', '흩어진 자료를 하나의 흐름으로 묶습니다', 9)
    cards(s, [
        ('1. AI가 제안합니다',
         '쌓인 뉴스·보도자료에서 "이건 하나의 흐름 같다"는 것을 이슈로 제안합니다.'),
        ('2. 운영자가 승인합니다',
         '제안을 그대로 쓰지 않습니다. 운영자가 확인해 승인한 것만 이슈가 됩니다.'),
        ('3. 연대기로 정리됩니다',
         '관련 뉴스·보도자료·법령·이해관계자가 시간 순으로 한 화면에 모입니다.'),
    ], Inches(1.95), cols=3, card_h=Inches(1.95), head_size=14.5)
    text(s, MARGIN, Inches(3.98), BODY_W, Inches(0.3),
         [('하나의 이슈에 모이는 것', 13, True, INK)])
    rows(s, [
        ('정부 보도자료 · 공고', '그 사안을 소관 부처가 어떻게 발표했는지'),
        ('정책 뉴스', '언제 무엇이 보도됐는지, 시간 순으로'),
        ('법령 · 법안', '관련 조문과 그 사이 발의된 개정안'),
        ('이해관계자', '누가 어느 자리에서 무슨 말을 했는지'),
    ], Inches(4.35), left_w=Inches(2.8), row_h=Inches(0.5), left_size=11.5, right_size=11)
    note(s, '"그 건 어떻게 흘러왔더라"를 다시 찾지 않으려고 만든 자리입니다.')

    # ── 10 국회 탭 ────────────────────────────────────────────
    s = blank(prs); head(s, '대시보드 ④ 국회', '2016년 20대 국회부터의 회의록을 사람 축으로도 봅니다', 10)
    rows(s, [
        ('법안', '발의 → 소관위 회부 → 상정 → 위원회 의결 → 법사위 → 본회의.\n'
                 '단계가 바뀌면 알리고, 개정안의 신구조문대비표도 조문 단위로 갈라 보여 줍니다'),
        ('입법예고', '국회·부처 입법예고를 모아 의견 마감일을 배지로 표시하고, 마감 3일 전에 다시 알립니다'),
        ('과방위 회의록', '2016년 20대 국회부터 상임위·국정감사 전문.\n'
                          '회의별·발언자별로 보거나, 원문 검색으로 찾습니다'),
        ('인물', '의원·정부 인사·기업 증인 243명.\n'
                 '쟁점별로 어떤 입장이었는지 발언 날짜를 근거로 붙여 정리합니다'),
    ], Inches(1.9), left_w=Inches(2.3), row_h=Inches(1.0))
    note(s, '입장 요약은 "질의했다·지적했다·답변했다"처럼 실제 행위로 씁니다 — 질의를 곧 신념으로 읽지 않습니다.')

    # ── 11 텔레그램 받기 (메일 3항) ───────────────────────────
    s = blank(prs); head(s, '텔레그램 — 받아 보기',
                         't.me/radio_policy_law_ai_bot  (정책 AI도우미)', 11, '메일 3항')
    cards(s, [
        ('① 링크 열기', '휴대폰에 텔레그램이 있으면\n링크를 누르는 것만으로 시작됩니다'),
        ('② /start', '받을 항목과 수신 시간대를 고르면 끝.\n설정은 /settings로 언제든 변경'),
    ], Inches(1.85), cols=2, card_h=Inches(1.3), head_size=15, body_size=11.5)
    text(s, MARGIN, Inches(3.4), BODY_W, Inches(0.35), [('받을 항목', 13.5, True, INK)])
    rows(s, [
        ('📡  모닝 브리핑', '시작 시각에 하루 한 번'),
        ('📡  주요 뉴스', '관심 분야 선택 가능'),
        ('🏛️  국회·법률 동향', '법안 발의·처리 경과, 국회·부처 입법예고, 과방위 회의록 요약'),
        ('📺  방미통위 동향', '회의 의사일정·의결 결과·보도자료'),
    ], Inches(3.85), left_w=Inches(2.9), row_h=Inches(0.62))
    note(s, '브리핑을 제외한 항목은 정해 두신 수신 시간대 안에서 새로 생기는 대로 전달됩니다.')

    # ── 12 텔레그램 물어보기 ──────────────────────────────────
    s = blank(prs); head(s, '텔레그램 — 물어보기',
                         '대화창에서 /를 누르면 메뉴에 뜹니다', 12, '메일 3항')
    cards(s, [
        ('/assem',
         '2023년 공공와이파이 관련 발언 찾아줘\n\n'
         '2016년부터 쌓인 과방위 발언에서 어느 의원이 언제 무슨 말을 했는지 원문 링크까지 찾아 줍니다.\n\n'
         '별도 승인 없이 바로 이용 가능'),
        ('/law',
         '전기통신사업법 19조\n\n'
         '조문 원문을 그대로 보여 드립니다.\n원문을 옮기는 것이라 AI가 지어낼 여지가 없습니다.\n\n'
         '별도 승인 없이 바로 이용 가능'),
    ], Inches(1.95), cols=2, card_h=Inches(2.85), head_size=17)
    rect(s, MARGIN, Inches(5.15), BODY_W, Inches(0.95), PANEL, LINE)
    text(s, MARGIN + Inches(0.32), Inches(5.38), BODY_W - Inches(0.64), Inches(0.6),
         [('주제로 묻는 법령 검색(/law 3G 종료 관련 법령)과 AI 자문(/ask)은 '
           '최초 1회 승인 후 이용하실 수 있습니다.', 12.5, False, INK)], line_spacing=1.3)

    # ── 13 다른 팀 + 계정 (메일 4·5항) ────────────────────────
    s = blank(prs); head(s, '다른 팀에서 쓰시려면', '현재 범위와 계획', 13, '메일 4·5항')
    text(s, MARGIN, Inches(1.68), BODY_W, Inches(0.6),
         [('이 시스템의 구조는 분야를 가리지 않습니다. 관심 키워드와 소관 법령, 담당 기관 목록만 바꾸면 '
           '어느 조직에서든 쓸 수 있습니다. 다만 지금은 그 목록이 전파·전기통신에 맞춰져 있습니다.',
           12.5, False, INK)], line_spacing=1.35)
    cards(s, [
        ('지금의 한계',
         '소관 법령이 다른 팀에서는 AI 자문과 뉴스 중요도 분류가 기대에 못 미칠 수 있습니다.\n'
         '특히 뉴스 중요도는 모든 이용자가 같은 값을 보는 구조라 지금은 조정을 제한해 두었습니다.'),
        ('앞으로의 계획',
         '팀별 설정(소관 법령·상임위·관점·뉴스 중요도)을 먼저 개발한 뒤 순차적으로 열어 드립니다.\n'
         '먼저 써 보실 팀은 소수 인원 파일럿으로 열고, 그 팀 법령을 반영하며 함께 맞춰 가겠습니다.'),
    ], Inches(2.5), cols=2, card_h=Inches(2.0), head_size=15)
    text(s, MARGIN, Inches(4.75), BODY_W, Inches(0.35), [('계정 안내', 13.5, True, INK)])
    rows(s, [
        ('계정 없이', '대시보드 열람 · 텔레그램 알림 · /assem 발언검색 · /law 조문 조회'),
        ('승인이 필요', 'AI 자문(대시보드·/ask), 주제로 묻는 법령 검색.\n'
                        '대시보드에서 가입 신청하시면 파일럿 팀부터 순차적으로 승인해 드립니다'),
    ], Inches(5.2), left_w=Inches(2.2), row_h=Inches(0.78), left_size=11.5, right_size=11)

    # ── 부록 A ────────────────────────────────────────────────
    s = blank(prs); head(s, '부록 A. 수집 기준', '이 범위로 매시간 검색합니다', None)
    text(s, MARGIN, Inches(1.72), BODY_W, Inches(0.3), [('정책 뉴스 수집 키워드', 13, True, INK)])
    rect(s, MARGIN, Inches(2.1), BODY_W, Inches(1.95), PANEL, LINE)
    text(s, MARGIN + Inches(0.3), Inches(2.35), BODY_W - Inches(0.6), Inches(1.5),
         [('주파수 · 5G/6G 주파수 · 전파정책 · 전파법 · 전파사용료 · 전자파 · 무선국 · 무선설비 · '
           '적합성평가 · 방송통신기자재 · WRC-27 · 6GHz · ITU · 위성통신', 11.5, False, INK),
          ('이동통신 · 기지국 · LTE · 3G · 이동통신 품질 · 통신품질 · 통신장애 · 이동통신 장비 · '
           '공공와이파이 · 지하철 와이파이', 11.5, False, INK),
          ('전기통신사업법 · 정보통신망법 · 위치정보법 · 통신정책 · 통신요금 · 알뜰폰 · 번호이동 · '
           '단말기유통 · 이용자보호 · 스팸문자 · 재난문자', 11.5, False, INK),
          ('사이버보안 · AI 기본법 · AI 규제 · 과기정통부 · 부처 인사(과기정통부 · 방미통위)',
           11.5, False, INK)], space_after=Pt(7), line_spacing=1.35)
    text(s, MARGIN, Inches(4.3), BODY_W, Inches(0.3), [('국회 법안 수집 키워드', 13, True, INK)])
    rect(s, MARGIN, Inches(4.68), BODY_W, Inches(1.15), PANEL, LINE)
    text(s, MARGIN + Inches(0.3), Inches(4.9), BODY_W - Inches(0.6), Inches(0.8),
         [('전파법 · 전기통신사업법 · 방송통신발전 · 정보통신망 · 주파수 · 전자파 · 무선국 · '
           '방송통신설비 · 적합성평가 · 이동통신단말 · 위성통신 · 기간통신 · 전파간섭', 11.5, False, INK),
          ('개인정보 · 인공지능 · 플랫폼 · 데이터 · 클라우드 · 메타버스   '
           '← 다른 팀 소재가 이미 상당 부분 들어와 있습니다', 11.5, True, INK)],
         space_after=Pt(7), line_spacing=1.35)
    note(s, '키워드는 넓게 긁고 관련 없는 것은 저장하지 않습니다 — 빠뜨리는 쪽보다 넉넉히 보는 쪽을 택했습니다.')

    # ── 부록 B ────────────────────────────────────────────────
    s = blank(prs); head(s, '부록 B. 지금 담고 있는 것', '2026년 9월 기준', None)
    rows(s, [
        ('과방위 회의록', '2016년 6월부터 393회 · 발언 6,741건 (상임위 · 국정감사 전문)'),
        ('인물', '243명 — 의원 125명 · 정부 인사와 기업 증인 118명. 쟁점별 입장 요약 포함'),
        ('국회 법안', '669건 — 발의부터 처리까지 단계 추적'),
        ('정책 뉴스', '누적 1만 1천여 건 · 매시간 수집 · 60일 보관'),
        ('법적 용어', '조문에 정의된 법률 용어 1,370개 — 원문 그대로'),
        ('법령 관계도', '법령·고시 노드 1,300여 개와 위임·인용 관계'),
    ], Inches(1.9), left_w=Inches(2.6), row_h=Inches(0.68))
    rect(s, MARGIN, Inches(5.95), BODY_W, Inches(0.95), PANEL, LINE)
    text(s, MARGIN + Inches(0.35), Inches(6.15), BODY_W - Inches(0.7), Inches(0.6),
         [('문의 · 기능 요청 · 자료 추가', 12.5, True, INK),
          ('SKT Comm.센터 기술정책팀   ·   radio-policy.gitlab.io   ·   '
           't.me/radio_policy_law_ai_bot', 11.5, False, MUTED)], space_after=Pt(5))

    out = '전파정책AI_소개_20260913.pptx'
    prs.save(out)
    print('저장: %s (%d장)' % (out, len(prs.slides._sldIdLst)))


if __name__ == '__main__':
    build()
